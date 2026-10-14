"""Build Obsidian notes from iCloud health, body scale, Scriptable, and food records."""

from __future__ import annotations

import json
import argparse
import os
import shutil
import sys
import calendar
import html
import csv
import subprocess
import time as time_module
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import openpyxl


TZ = timezone(timedelta(hours=8))
VAULT = HEALTH_SOURCE = SCRIPTABLE = ARCHIVE = BOOK = None
SCALE_REPO = SCALE_CSV = SCALE_API = None
SCALE_SOURCE_LABEL = "体脂秤"
SCHEDULE_TIME = time(9, 30)


def configure(config_path):
    """Load local paths. The real config is intentionally excluded from Git."""
    global VAULT, HEALTH_SOURCE, SCRIPTABLE, ARCHIVE, BOOK
    global SCALE_REPO, SCALE_CSV, SCALE_API, SCALE_SOURCE_LABEL, SCHEDULE_TIME
    path = Path(config_path).expanduser().resolve()
    config = read_json(path)
    if not isinstance(config, dict):
        raise ValueError("配置文件必须是 JSON 对象")

    def local_path(key, required=True):
        value = config.get(key)
        if not isinstance(value, str) or not value.strip():
            if required:
                raise ValueError(f"请在配置文件中填写 {key}")
            return None
        candidate = Path(os.path.expandvars(value)).expanduser()
        return (candidate if candidate.is_absolute() else path.parent / candidate).resolve()

    VAULT = local_path("vault_path")
    HEALTH_SOURCE = local_path("health_export_dir")
    SCRIPTABLE = local_path("scriptable_dir")
    BOOK = local_path("diet_workbook")
    ARCHIVE = local_path("archive_dir", required=False)
    if ARCHIVE is None:
        base = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
        ARCHIVE = base / "HealthObsidianSync" / "archive"
    SCALE_CSV = local_path("scale_csv", required=False)
    SCALE_REPO = local_path("scale_repo_path", required=False)
    github_repo = config.get("scale_github_repo", "")
    if github_repo:
        if not isinstance(github_repo, str) or github_repo.count("/") != 1 or not SCALE_REPO or not SCALE_CSV:
            raise ValueError("scale_github_repo 需要 owner/repo、scale_repo_path 和 scale_csv")
        SCALE_API = "https://api.github.com/repos/" + github_repo + "/actions"
    else:
        SCALE_API = None
    SCALE_SOURCE_LABEL = str(config.get("scale_source_label") or "体脂秤")
    schedule = str(config.get("schedule_time") or "09:30")
    try:
        SCHEDULE_TIME = time.fromisoformat(schedule)
    except ValueError:
        raise ValueError("schedule_time 应为 HH:MM") from None


def scale_status(synced):
    if synced is None:
        return "未启用云端采集，显示本地记录"
    return "本次已更新" if synced else "本次未更新，显示已有记录"
BEGIN = "<!-- 自动生成内容开始：请把个人笔记写在页面末尾的“个人笔记”中 -->"
END = "<!-- 自动生成内容结束 -->"
STATUS = {0: "待记录", 1: "已完成", 2: "完成一部分", 3: "一点没干"}
METRIC_NAMES = {
    "Walking Speed": "步行速度", "Walking Step Length": "步长", "Flights Climbed": "爬楼层数",
    "Walking Heart Rate": "步行心率", "Heart Rate": "心率", "Resting Heart Rate": "静息心率",
    "Active Energy": "活动能量", "Basal Energy": "基础能量", "Walking Asymmetry": "步态不对称",
    "Walking + Running Distance": "步行与跑步距离", "Steps": "步数",
    "Respiratory Rate": "呼吸频率", "Cycling Distance": "骑行距离",
    "Heart Rate Variability": "心率变异性",
}
SLEEP_LABELS = {"asleep": "睡着", "asleepUnspecified": "睡着（未分期）", "awake": "清醒", "unknown": "未知"}


def local_time(value):
    if not value:
        return None
    return datetime.fromisoformat(str(value).replace("Z", "+00:00")).astimezone(TZ)


def cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip().replace("\n", "<br>").replace("|", "\\|")


def as_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10]).isoformat()
        except ValueError:
            return None
    return None


def copy_new_files(source, pattern, target):
    target.mkdir(parents=True, exist_ok=True)
    if not source.is_dir():
        return
    for path in source.glob(pattern):
        if not path.is_file():
            continue
        copy = target / path.name
        try:
            if not copy.exists() or path.stat().st_size != copy.stat().st_size or path.stat().st_mtime_ns > copy.stat().st_mtime_ns:
                shutil.copy2(path, copy)
        except OSError as exc:
            print(f"暂时无法读取同步文件：{path.name}: {exc}", file=sys.stderr)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def github_headers():
    env = os.environ.copy()
    env["GCM_INTERACTIVE"] = "never"
    env["GIT_TERMINAL_PROMPT"] = "0"
    result = subprocess.run(
        ["git", "credential", "fill"],
        input="protocol=https\nhost=github.com\n\n",
        text=True, capture_output=True, timeout=15, env=env,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    values = dict(line.split("=", 1) for line in result.stdout.splitlines() if "=" in line)
    if result.returncode or not values.get("password"):
        raise RuntimeError("这台电脑的 GitHub 授权不可用")
    return {
        "Authorization": "Bearer " + values["password"],
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }


def github_api(path, headers, payload=None):
    request = urllib.request.Request(
        SCALE_API + path,
        data=json.dumps(payload).encode("utf-8") if payload is not None else None,
        headers={**headers, "Content-Type": "application/json"},
        method="POST" if payload is not None else "GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            body = response.read()
            return json.loads(body) if body else {}
    except urllib.error.HTTPError as exc:
        raise RuntimeError(f"GitHub 同步接口返回 HTTP {exc.code}") from None
    except urllib.error.URLError:
        raise RuntimeError("暂时无法连接 GitHub 同步接口") from None


def sync_scale():
    """Refresh the private scale CSV before building local Obsidian pages."""
    if SCALE_API is None:
        return None
    try:
        headers = github_headers()
        before = github_api("/workflows/sync.yml/runs?event=workflow_dispatch&per_page=5", headers)
        previous_ids = {run["id"] for run in before.get("workflow_runs", [])}
        github_api("/workflows/sync.yml/dispatches", headers, {"ref": "main", "inputs": {"quantity": "100"}})
        deadline = time_module.monotonic() + 150
        run_id = None
        while time_module.monotonic() < deadline:
            if run_id is None:
                runs = github_api("/workflows/sync.yml/runs?event=workflow_dispatch&per_page=5", headers)
                newer = [run for run in runs.get("workflow_runs", []) if run["id"] not in previous_ids]
                if newer:
                    run_id = newer[0]["id"]
            if run_id is not None:
                run = github_api(f"/runs/{run_id}", headers)
                if run.get("status") == "completed":
                    if run.get("conclusion") != "success":
                        raise RuntimeError("体脂秤云端同步未成功")
                    break
            time_module.sleep(5)
        else:
            raise RuntimeError("等待体脂秤云端同步超时")
        env = os.environ.copy()
        env["GCM_INTERACTIVE"] = "never"
        env["GIT_TERMINAL_PROMPT"] = "0"
        pulled = subprocess.run(
            ["git", "-C", str(SCALE_REPO), "pull", "--ff-only"],
            text=True, capture_output=True, timeout=40, env=env,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if pulled.returncode:
            raise RuntimeError("无法更新本地体脂秤记录")
        return True
    except (OSError, RuntimeError, subprocess.TimeoutExpired, ValueError) as exc:
        print(f"体脂秤本次使用已有记录：{exc}", file=sys.stderr)
        return False


def scale_data():
    if SCALE_CSV is None or not SCALE_CSV.exists():
        return []
    with SCALE_CSV.open("r", newline="", encoding="utf-8-sig") as file:
        records = {row["historyId"]: row for row in csv.DictReader(file) if row.get("historyId")}
    return sorted(records.values(), key=lambda row: int(row["recordDate"]))


def health_data():
    source_dir = ARCHIVE / "health"
    copy_new_files(HEALTH_SOURCE, "*.json", source_dir)
    files = sorted(source_dir.glob("*.json"), key=lambda p: (p.stat().st_mtime_ns, p.name))
    metrics, categories, workouts = {}, {}, {}
    used = []
    for path in files:
        try:
            data = read_json(path)
        except (OSError, ValueError) as exc:
            print(f"跳过无法读取的健康文件：{path.name}: {exc}", file=sys.stderr)
            continue
        if not isinstance(data, dict) or "metrics" not in data:
            continue
        used.append(path.name)
        for metric in data.get("metrics", []):
            identity = metric.get("id") or metric.get("display_name")
            for point in metric.get("data_points", []):
                key = (identity, point.get("start_date") or point.get("timestamp"), point.get("end_date"))
                metrics[key] = (metric, point)
        for category in data.get("category_metrics", []):
            identity = category.get("id") or category.get("display_name")
            for point in category.get("data_points", []):
                source = point.get("source") or {}
                key = (identity, point.get("start_date") or point.get("timestamp"), point.get("end_date"), source.get("bundle_identifier") or source.get("name"))
                categories[key] = (category, point)
        for workout in data.get("workouts", []):
            key = workout.get("uuid") or workout.get("id") or json.dumps(workout, ensure_ascii=False, sort_keys=True)
            workouts[key] = workout
    return list(metrics.values()), list(categories.values()), list(workouts.values()), used


def scriptable_data():
    copy_new_files(SCRIPTABLE / "DailyBlocks", "*.json", ARCHIVE / "daily")
    copy_new_files(SCRIPTABLE / "TrainingLogs", "*.json", ARCHIVE / "training")
    daily = {}
    for path in sorted((ARCHIVE / "daily").glob("*.json")):
        try:
            data = read_json(path)
            day = as_date(data.get("date"))
            if day and isinstance(data.get("tasks"), list):
                daily[day] = data["tasks"]
        except (OSError, ValueError, AttributeError) as exc:
            print(f"跳过无法读取的每日事项：{path.name}: {exc}", file=sys.stderr)
    training = {}
    for path in sorted((ARCHIVE / "training").glob("*.json"), key=lambda p: p.stat().st_mtime_ns):
        try:
            data = read_json(path)
            for day, records in data.get("records", {}).items():
                if as_date(day) and isinstance(records, list):
                    training[day] = records
        except (OSError, ValueError, AttributeError) as exc:
            print(f"跳过无法读取的健身记录：{path.name}: {exc}", file=sys.stderr)
    return daily, training


def workbook_data():
    book = BOOK
    if not book.exists():
        raise FileNotFoundError(f"找不到饮食 Excel：{BOOK}")
    wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
    try:
        ws = wb["饮食与运动记录"] if "饮食与运动记录" in wb else wb.active
        foods, exercises, notes = [], [], []
        current_day, current_period = None, ""
        for row in ws.iter_rows(min_row=9, max_col=11, values_only=True):
            day = as_date(row[0])
            if day:
                current_day, current_period = day, ""
            if row[1] is not None:
                current_period = cell(row[1])
            if not current_day:
                continue
            if any(x is not None and str(x).strip() for x in row[2:6]):
                foods.append((current_day, current_period, *[cell(x) for x in row[2:6]]))
            if any(x is not None and str(x).strip() for x in row[6:10]):
                exercises.append((current_day, current_period, *[cell(x) for x in row[6:10]]))
            if row[10] is not None and str(row[10]).strip():
                notes.append((current_day, str(row[10]).strip()))
        profile = []
        for row in ws.iter_rows(min_row=4, max_row=5, max_col=10, values_only=True):
            profile.extend(cell(x) for x in row if x is not None and str(x).strip())
        return foods, exercises, notes, profile, book
    finally:
        wb.close()


def table(headers, rows):
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    lines.extend("| " + " | ".join(cell(x) for x in row) + " |" for row in rows)
    return "\n".join(lines) if rows else "暂无记录。"


def stamp():
    return datetime.now(TZ).strftime("%Y-%m-%d %H:%M")


def intro(title):
    return f"# {title}\n\n<div class=\"health-updated\">更新于 {stamp()} · 北京时间 · 计划每天 {SCHEDULE_TIME:%H:%M} 整理</div>\n"


def group_by_month(rows, date_index=0):
    months = defaultdict(list)
    for row in rows:
        months[str(row[date_index])[:7]].append(row)
    return months


def group_by_day(rows, date_index=0):
    days = defaultdict(list)
    for row in rows:
        days[str(row[date_index])].append(row)
    return days


def folded_callout(title, body):
    """Obsidian's native foldable callout keeps long source tables available."""
    return "\n".join([f"> [!note]- {title}", ">"] + ["> " + line if line else ">" for line in body.splitlines()])


def chart_block(kind, title, values, y_title, limit=30):
    if not values:
        return "暂无可绘制的数据。"
    valid_days = sorted(date.fromisoformat(day) for day in values)
    end = valid_days[-1]
    start = max(valid_days[0], end - timedelta(days=limit - 1))
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    labels = [day.strftime("%m/%d") for day in days]
    data = [round(float(values[day.isoformat()]), 2) if day.isoformat() in values else None for day in days]
    fence = chr(96) * 3
    return "\n".join([
        fence + "chart", f"type: {kind}", f"labels: {json.dumps(labels, ensure_ascii=False)}",
        "series:", f"  - title: {json.dumps(title, ensure_ascii=False)}", f"    data: {json.dumps(data, ensure_ascii=False)}",
        f"yTitle: {json.dumps(y_title, ensure_ascii=False)}", "legend: false", "spanGaps: false",
        "beginAtZero: true" if kind == "bar" else "tension: 0.25", fence,
    ])


def category_chart(title, counts, kind="doughnut"):
    """Chart only categories that actually occur in the source records."""
    items = [(label, count) for label, count in counts.items() if count]
    if not items:
        return "暂无可绘制的数据。"
    fence = chr(96) * 3
    return "\n".join([
        fence + "chart", f"type: {kind}",
        f"labels: {json.dumps([label for label, _ in items], ensure_ascii=False)}",
        "series:", f"  - title: {json.dumps(title, ensure_ascii=False)}",
        f"    data: {json.dumps([count for _, count in items], ensure_ascii=False)}",
        "legend: true", "width: 65%", fence,
    ])


def nav_callout():
    return "\n".join([
        "> [!health-nav] 浏览专题",
        "> - [[健康指标|📈 健康指标]]  ·  步数、心率与全部导出指标",
        "> - [[睡眠记录|🌙 睡眠记录]]  ·  时长与原始状态",
        "> - [[体脂秤|⚖️ 体脂秤]]  ·  体重与身体成分",
        "> - [[饮食记录|🍽️ 饮食记录]]  ·  Excel 原始条目",
        "> - [[健身与运动|🏃 健身与运动]]  ·  三处运动来源",
        "> - [[每日事项|✅ 每日事项]]  ·  计划与完成情况",
        "> - [[饮食训练日历|🗓️ 饮食训练日历]]  ·  按日期进入记录",
    ])


def stat_cards(cards):
    blocks = []
    for label, value, detail in cards:
        blocks.append(
            '<div class="health-stat"><span class="health-stat-label">'
            + html.escape(str(label)) + '</span><strong>' + html.escape(str(value))
            + '</strong><small>' + html.escape(str(detail)) + '</small></div>'
        )
    return '<div class="health-stat-grid">\n' + "\n".join(blocks) + '\n</div>'


def latest_value(series):
    if not series:
        return None, None
    day = max(series)
    return day, series[day]


def scale_time(record):
    return datetime.fromtimestamp(int(record["recordDate"]) / 1000, TZ)


def scale_number(record, field):
    try:
        return float(record[field])
    except (KeyError, TypeError, ValueError):
        return None


def scale_series(records, field):
    by_day = {}
    for record in records:
        value = scale_number(record, field)
        if value is not None:
            by_day[scale_time(record).date().isoformat()] = value
    return by_day


def scale_page(records, synced):
    latest = records[-1] if records else None
    latest_time = scale_time(latest).strftime("%Y-%m-%d %H:%M") if latest else "暂无记录"
    weight = scale_number(latest, "weight") if latest else None
    fat = scale_number(latest, "bodyFat") if latest else None
    bmi = scale_number(latest, "bmi") if latest else None
    muscle = scale_number(latest, "skeletalMuscle") if latest else None
    water = scale_number(latest, "moisture") if latest else None
    visceral = scale_number(latest, "visceralFat") if latest else None
    lines = [
        intro("体脂秤"),
        stat_cards([
            ("最新体重", f"{weight:.1f} kg" if weight is not None else "暂无", latest_time),
            ("最新体脂率", f"{fat:.2f}%" if fat is not None else "暂无", latest_time),
            ("最新 BMI", f"{bmi:.2f}" if bmi is not None else "暂无", latest_time),
            ("骨骼肌率", f"{muscle:.2f}%" if muscle is not None else "暂无", f"{SCALE_SOURCE_LABEL}原值"),
            ("体水分率", f"{water:.2f}%" if water is not None else "暂无", f"{SCALE_SOURCE_LABEL}原值"),
            ("内脏脂肪", f"{visceral:g}" if visceral is not None else "暂无", f"{SCALE_SOURCE_LABEL}原值"),
            ("历史测量", f"{len(records)} 次", "按记录 ID 去重"),
        ]),
        f"体脂秤同步：{scale_status(synced)}；最新测量时间：{latest_time}。",
        f"每天的趋势点取当天最后一次测量。原始数值来自{SCALE_SOURCE_LABEL}；这里只展示记录，不据此作健康诊断。",
        "## 最近 30 天趋势",
        "### 体重",
        chart_block("line", "体重", scale_series(records, "weight"), "kg"),
        "### 体脂率",
        chart_block("line", "体脂率", scale_series(records, "bodyFat"), "%"),
        "### 骨骼肌率",
        chart_block("line", "骨骼肌率", scale_series(records, "skeletalMuscle"), "%"),
        "### 体水分率",
        chart_block("line", "体水分率", scale_series(records, "moisture"), "%"),
        "## 全部测量",
        table(
            ["测量时间", "体重 (kg)", "体脂率 (%)", "BMI"],
            [
                (
                    scale_time(record).strftime("%Y-%m-%d %H:%M"),
                    f"{scale_number(record, 'weight'):.2f}" if scale_number(record, "weight") is not None else "",
                    f"{scale_number(record, 'bodyFat'):.2f}" if scale_number(record, "bodyFat") is not None else "",
                    f"{scale_number(record, 'bmi'):.2f}" if scale_number(record, "bmi") is not None else "",
                )
                for record in reversed(records)
            ],
        ),
        "## 完整原始指标",
        "展开任一测量可查看 CSV 中保存的全部字段。除上方明确标注的单位外，下表保留服务端原值。",
    ]
    for record in reversed(records):
        measured = scale_time(record).strftime("%Y-%m-%d %H:%M")
        lines.append(folded_callout(measured, table(["原始字段", "原始数值"], record.items())))
    return "\n\n".join(lines)


def metrics_page(metrics, categories, used):
    rows = []
    index = defaultdict(dict)
    for metric, point in metrics:
        dt = local_time(point.get("timestamp") or point.get("start_date"))
        if dt is None:
            continue
        name = metric.get("display_name") or metric.get("id") or "未命名指标"
        day = dt.date().isoformat()
        value = point.get("value", "")
        unit = point.get("unit") or metric.get("unit") or ""
        note = ""
        if name in ("Walking + Running Distance", "Cycling Distance"):
            note = "距离数值或单位待核对；可能混入非步行移动"
            unit = f"{unit}（导出标注）" if unit else "导出单位缺失"
        elif name == "Walking Asymmetry":
            note = "按导出原值保留，未换算百分比"
        rows.append((day, METRIC_NAMES.get(name, name), value, unit, note, name))
        if isinstance(value, (int, float)):
            index[name][day] = value
    rows.sort(key=lambda r: (r[0], r[5]), reverse=True)
    days = {r[0] for r in rows}
    recent = sorted(days, reverse=True)[:7]
    summary = []
    for day in recent:
        steps = index.get("Steps", {}).get(day)
        resting = index.get("Resting Heart Rate", {}).get(day)
        energy = index.get("Active Energy", {}).get(day)
        summary.append((day, f"{steps:,.0f}" if steps is not None else "未记录", f"{resting:.0f}" if resting is not None else "未记录", f"{energy:.0f}" if energy is not None else "未记录"))
    lines = [
        intro("健康指标"),
        stat_cards([
            ("最新步数", f"约 {latest_value(index.get('Steps', {}))[1]:,.0f}" if index.get("Steps") else "暂无", latest_value(index.get("Steps", {}))[0] or "暂无日期"),
            ("最新静息心率", f"{latest_value(index.get('Resting Heart Rate', {}))[1]:.0f} 次/分" if index.get("Resting Heart Rate") else "暂无", latest_value(index.get("Resting Heart Rate", {}))[0] or "暂无日期"),
            ("数值记录", str(len(rows)), f"{len(used)} 份健康导出"),
        ]),
        f"来自 {len(used)} 份导出，共 {len(rows)} 条数值记录。下方摘要便于查看，所有原始数值仍保留在折叠明细中。",
        "> [!warning] 距离记录待核对\n> 请核对导出文件的距离数值、单位和运动来源；部分来源可能混入非步行移动。本页保留原值，首页图表不使用距离数据。",
        "## 有来源的健康趋势",
        "### 每日步数",
        chart_block("bar", "步数", index.get("Steps", {}), "步"),
        "### 静息心率",
        chart_block("line", "静息心率", index.get("Resting Heart Rate", {}), "次/分"),
        "### 活动能量",
        chart_block("bar", "活动能量", index.get("Active Energy", {}), "千卡"),
        "### 心率变异性",
        chart_block("line", "心率变异性", index.get("Heart Rate Variability", {}), "毫秒"),
        "## 最近 7 个有记录的日期",
        table(["日期", "步数（约）", "静息心率（次/分）", "活动能量（千卡）"], summary),
        "## 完整指标明细",
    ]
    for month, items in sorted(group_by_month(rows).items(), reverse=True):
        lines += [folded_callout(f"{month} · {len(items)} 条原始数值", table(["日期", "指标", "原始数值", "原始单位", "备注"], [r[:5] for r in items])), ""]
    if not rows:
        lines.append("暂无记录。")
    other = [(category, point) for category, point in categories if (category.get("id") or category.get("display_name")) != "Sleep Analysis"]
    if other:
        lines.append("## 其他分类记录\n")
        for category, point in other:
            lines.append("```json\n" + json.dumps({"type": category.get("display_name"), "record": point}, ensure_ascii=False, indent=2) + "\n```\n")
    return "\n\n".join(lines), len(rows), days, index


def sleep_day(dt):
    """Attribute evening-to-morning sleep to the morning date (18:00 boundary)."""
    return (dt.date() + timedelta(days=1) if dt.hour >= 18 else dt.date()).isoformat()


def sleep_parts(start, end):
    cursor = start
    while cursor < end:
        boundary = datetime.combine(cursor.date(), time(18), tzinfo=TZ)
        if cursor >= boundary:
            boundary += timedelta(days=1)
        stop = min(end, boundary)
        yield sleep_day(cursor), cursor, stop
        cursor = stop


def merged_hours(intervals):
    """Prevent overlapping exports or sources from counting the same minute twice."""
    total = 0.0
    last_end = None
    for start, end in sorted(intervals):
        if last_end is None or start > last_end:
            total += (end - start).total_seconds()
            last_end = end
        elif end > last_end:
            total += (end - last_end).total_seconds()
            last_end = end
    return total / 3600


def sleep_timeline_svg(segments):
    """A fixed night-axis view of the actual exported sleep states."""
    days = sorted(segments, reverse=True)[:7]
    if not days:
        return None
    width, left, plot = 960, 142, 770
    height = 104 + 44 * len(days)
    palette = {"asleep": "#50c5aa", "asleepUnspecified": "#50c5aa", "awake": "#ffab78", "unknown": "#8392a9"}
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" role="img" aria-label="最近七个睡眠日期的原始状态时间轴">',
           '<rect width="100%" height="100%" rx="18" fill="#182233"/>',
           '<text x="20" y="31" fill="#f3f7fb" font-size="17" font-family="sans-serif" font-weight="700">睡眠状态时间轴</text>',
           '<text x="20" y="52" fill="#a8b7c9" font-size="12" font-family="sans-serif">每行从前一天 18:00 到当天 12:00 · 北京时间</text>']
    for hour in (18, 21, 0, 3, 6, 9, 12):
        offset = (hour - 18) % 24
        if hour == 12:
            offset = 18
        x = left + plot * offset / 18
        out.append(f'<line x1="{x:.1f}" y1="68" x2="{x:.1f}" y2="{height-35}" stroke="#314154" stroke-width="1"/>')
        out.append(f'<text x="{x:.1f}" y="70" text-anchor="middle" fill="#a8b7c9" font-size="11" font-family="sans-serif">{hour:02d}:00</text>')
    for i, day in enumerate(days):
        d = date.fromisoformat(day)
        window_start = datetime.combine(d - timedelta(days=1), time(18), tzinfo=TZ)
        window_end = window_start + timedelta(hours=18)
        y = 85 + i * 44
        out.append(f'<text x="20" y="{y+17}" fill="#f3f7fb" font-size="13" font-family="sans-serif">{html.escape(day)}</text>')
        out.append(f'<rect x="{left}" y="{y}" width="{plot}" height="23" rx="5" fill="#28384b"/>')
        for start, end, label in sorted(segments[day]):
            a, b = max(start, window_start), min(end, window_end)
            if b <= a:
                continue
            x = left + plot * (a - window_start).total_seconds() / 64800
            w = max(1, plot * (b - a).total_seconds() / 64800)
            fill = palette.get(label, "#8392a9")
            out.append(f'<rect x="{x:.1f}" y="{y}" width="{w:.1f}" height="23" rx="2" fill="{fill}"/>')
    legend_y = height - 13
    for x, color, label in ((20, "#50c5aa", "睡着（未细分）"), (204, "#ffab78", "清醒"), (304, "#8392a9", "未知")):
        out.append(f'<rect x="{x}" y="{legend_y-9}" width="12" height="12" rx="2" fill="{color}"/>')
        out.append(f'<text x="{x+18}" y="{legend_y+1}" fill="#d6e2ed" font-size="12" font-family="sans-serif">{label}</text>')
    out.append('</svg>')
    asset = VAULT / "_assets" / "sleep-timeline.svg"
    asset.parent.mkdir(parents=True, exist_ok=True)
    asset.write_text("\n".join(out), encoding="utf-8")
    return "![[sleep-timeline.svg]]"


def sleep_page(categories):
    rows = []
    intervals = defaultdict(list)
    segments = defaultdict(list)
    for category, point in categories:
        if (category.get("id") or category.get("display_name")) != "Sleep Analysis":
            continue
        start, end = local_time(point.get("start_date")), local_time(point.get("end_date"))
        if not start or not end or end <= start:
            continue
        label = point.get("label") or str(point.get("value", ""))
        minutes = (end - start).total_seconds() / 60
        rows.append((sleep_day(start), start.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S"), SLEEP_LABELS.get(label, label), f"{minutes:.1f}", (point.get("source") or {}).get("name", "")))
        for day, part_start, part_end in sleep_parts(start, end):
            segments[day].append((part_start, part_end, label))
            if label in ("asleep", "asleepUnspecified"):
                intervals[day].append((part_start, part_end))
    asleep = {day: round(merged_hours(parts), 3) for day, parts in intervals.items()}
    rows.sort(key=lambda r: r[1], reverse=True)
    summary = [(day, f"{hours:.2f}") for day, hours in sorted(asleep.items(), reverse=True)]
    latest_day, latest_hours = latest_value(asleep)
    lines = [
        intro("睡眠记录"),
        stat_cards([("最近一晚已标记睡着", f"{latest_hours:.2f} 小时" if latest_hours is not None else "暂无", latest_day or "暂无日期"), ("原始片段", str(len(rows)), "含睡着、清醒和未知")]),
        "以北京时间 18:00 作为夜间分界，夜晚记录归到次日；重叠的“睡着”片段只计一次。未知和清醒不算入睡着时长。本模板不推断深睡／REM 分期。",
        "## 最近七晚状态",
        sleep_timeline_svg(segments) or "暂无可绘制的片段。",
        "## 最近 30 天趋势",
        chart_block("bar", "已标记睡着时长", asleep, "小时"),
        "## 按夜晚归属的记录",
        table(["日期（醒来日）", "已标记睡着（小时）"], summary),
        "## 原始睡眠片段",
    ]
    for month, items in sorted(group_by_month(rows).items(), reverse=True):
        lines += [folded_callout(f"{month} · {len(items)} 条片段", table(["归属日期", "开始", "结束", "状态", "分钟", "来源"], items)), ""]
    if not rows:
        lines.append("暂无记录。")
    return "\n\n".join(lines), len(rows), asleep


def workout_day(workout):
    for key in ("start_date", "startDate", "start", "timestamp"):
        if workout.get(key):
            try:
                return local_time(workout[key]).date().isoformat()
            except (ValueError, TypeError):
                return as_date(workout[key])
    return None


def fitness_page(training, exercises, workouts):
    by_day = group_by_day(exercises)
    workout_days = defaultdict(list)
    undated = []
    for workout in workouts:
        day = workout_day(workout)
        if day:
            workout_days[day].append(workout)
        else:
            undated.append(workout)
    days = set(training) | set(by_day) | set(workout_days)
    monthly = Counter(day[:7] for day in days)
    lines = [
        intro("健身与运动"),
        stat_cards([("有记录的日期", str(len(days)), "三处来源合并去重"), ("饮食表运动", str(len(exercises)), "原始运动条目"), ("Scriptable 健身", str(sum(len(x) for x in training.values())), "原始健身记录"), ("健康体能训练", str(len(workouts)), "健康导出记录")]),
        "记录来自 Scriptable、饮食表和 Apple 健康。相同日期的来源并列展示，不把它们相加为一次训练。",
        "[[饮食训练日历|打开饮食训练日历]]",
        "## 每月有记录的训练日",
        category_chart("有记录的日期", dict(sorted(monthly.items())), kind="bar"),
        "## 按日期查看",
    ]
    for day in sorted(days, reverse=True):
        lines.append(f"## {day}")
        if training.get(day):
            lines.append("**Scriptable**：" + "、".join(cell(x) for x in training[day]))
        if by_day.get(day):
            lines += ["**饮食表中的运动**", table(["时段", "项目", "时长", "阻力／坡度／重量", "配速／组数与次数"], [row[1:] for row in by_day[day]])]
        for workout in workout_days.get(day, []):
            lines += ["**Apple 健康体能训练原始记录**", "```json\n" + json.dumps(workout, ensure_ascii=False, indent=2) + "\n```"]
    if undated:
        lines.append("## 日期未识别的 Apple 健康体能训练")
        for workout in undated:
            lines.append("```json\n" + json.dumps(workout, ensure_ascii=False, indent=2) + "\n```")
    if not workouts:
        lines.append("> [!info] Apple 健康体能训练\n> 当前导出中没有此类记录。")
    return "\n\n".join(lines), days


def food_page(foods, notes, profile, book):
    by_day = group_by_day(foods)
    latest_day = max(by_day) if by_day else "暂无"
    period_counts = Counter(row[1] or "未填写时段" for row in foods)
    lines = [
        intro("饮食记录"),
        stat_cards([("已记录天数", str(len(by_day)), "Excel 饮食表"), ("食物与饮品", str(len(foods)), f"最新 {latest_day}"), ("有时段标注", str(sum(count for period, count in period_counts.items() if period != "未填写时段")), "按 Excel 原列统计")]),
        f"来源：{book.name}。按原表逐条保留食物、饮品、数量和备注；没有根据名称推算热量或营养素。",
        "[[饮食训练日历|打开饮食训练日历]]",
        "## 记录概览",
        "### 最近 30 天每日条目数",
        chart_block("bar", "饮食条目", {day: len(items) for day, items in by_day.items()}, "条"),
        "### 按原表时段统计",
        category_chart("饮食条目", dict(period_counts)),
        "## 按日期查看原始内容",
    ]
    last_month = None
    for day, items in sorted(by_day.items(), reverse=True):
        month = day[:7]
        if month != last_month:
            lines.append(f"## {month}")
            last_month = month
        lines += [f"### {day}", table(["时段", "食物／饮品", "数量", "重量／容量", "备注"], [row[1:] for row in items])]
    if not foods:
        lines.append("暂无记录。")
    if notes:
        lines.append("## 表格中的营养评估原文")
        for day, note in sorted(notes, reverse=True):
            lines.append(folded_callout(f"{day} · 查看 Excel 原始备注", note))
    if profile:
        lines += ["## 表格中的个人资料与目标原文", "这些内容来自 Excel 表头，未重新计算。", folded_callout("查看原文", "\n".join(f"- {item}" for item in profile))]
    return "\n\n".join(lines)


def minute_string(value):
    try:
        minutes = int(value)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    except (ValueError, TypeError):
        return cell(value)


def tasks_page(daily):
    latest_day = max(daily) if daily else None
    latest_tasks = daily.get(latest_day, []) if latest_day else []
    done = sum(item.get("status") == 1 for item in latest_tasks)
    pending = sum(item.get("status") == 0 for item in latest_tasks)
    all_tasks = [task for tasks in daily.values() for task in tasks]
    status_counts = Counter(STATUS.get(task.get("status"), "其他状态") for task in all_tasks)
    lines = [
        intro("每日事项"),
        stat_cards([("最近记录", latest_day or "暂无", f"{len(latest_tasks)} 项计划"), ("已完成", str(done), "最近记录日"), ("待记录", str(pending), "最近记录日"), ("全部事项", str(len(all_tasks)), "Scriptable 原始条目")]),
        "## 全部记录的状态",
        category_chart("事项状态", dict(status_counts)),
        "## 按日期查看",
    ]
    for day, tasks in sorted(daily.items(), reverse=True):
        rows = [(minute_string(item.get("start")), minute_string(item.get("end")), item.get("goal", ""), STATUS.get(item.get("status"), cell(item.get("status"))), item.get("actual", "")) for item in sorted(tasks, key=lambda x: x.get("start", 0))]
        lines += [f"## {day}", table(["开始", "结束", "计划", "状态", "实际情况"], rows), ""]
    if not daily:
        lines.append("暂无记录。")
    return "\n".join(lines)


def calendar_table(month, food_days, training_days):
    weeks = []
    for week in calendar.Calendar(firstweekday=0).monthdatescalendar(month.year, month.month):
        cells = []
        for day in week:
            if day.month != month.month:
                cells.append("")
                continue
            iso = day.isoformat()
            count = len(food_days.get(iso, []))
            level = 3 if count >= 6 else 2 if count >= 3 else 1 if count else 0
            marker = f'<span class="health-day-marker health-level-{level}"></span>'
            target = "饮食记录" if count else "健身与运动" if iso in training_days else None
            day_label = f"[[{target}#{iso}|{day.day}]]" if target else f"**{day.day}**"
            parts = [marker + day_label]
            if count:
                parts.append(f"[[饮食记录#{iso}|饮食 {count}]]")
            if iso in training_days:
                parts.append(f"[[健身与运动#{iso}|训练]]")
            cells.append("<br>".join(parts))
        weeks.append(cells)
    return table(["周一", "周二", "周三", "周四", "周五", "周六", "周日"], weeks)


def calendar_page(foods, training_days):
    food_days = group_by_day(foods)
    recorded = set(food_days) | set(training_days)
    lines = [
        intro("饮食训练日历"),
        "点日期数字可进入当天记录；同一天有饮食和训练时，格内分别提供入口。颜色只表示饮食表的条目数量，不表示摄入量或训练强度。",
        '<div class="health-calendar-legend"><span class="health-level-0">无饮食条目</span><span class="health-level-1">1–2 条</span><span class="health-level-2">3–5 条</span><span class="health-level-3">6 条及以上</span></div>',
    ]
    if not recorded:
        return "\n\n".join(lines + ["暂无记录。"])
    first = date.fromisoformat(min(recorded)).replace(day=1)
    last = max(date.fromisoformat(max(recorded)).replace(day=1), datetime.now(TZ).date().replace(day=1))
    months = []
    month = first
    while month <= last:
        months.append(month)
        month = date(month.year + (month.month == 12), month.month % 12 + 1, 1)
    for month in reversed(months):
        lines += [f"## {month:%Y-%m}", calendar_table(month, food_days, training_days)]
    return "\n\n".join(lines)


def overview_page(metric_count, days, health_index, sleep_hours, sleep_count, foods, exercises, daily, training, workouts, health_files, book, scale_records, scale_synced):
    step_day, steps = latest_value(health_index.get("Steps", {}))
    heart_day, resting = latest_value(health_index.get("Resting Heart Rate", {}))
    sleep_day, hours = latest_value(sleep_hours)
    food_day = max((row[0] for row in foods), default=None)
    latest_scale = scale_records[-1] if scale_records else None
    scale_day = scale_time(latest_scale).date().isoformat() if latest_scale else "暂无日期"
    weight = scale_number(latest_scale, "weight") if latest_scale else None
    fat = scale_number(latest_scale, "bodyFat") if latest_scale else None
    food_days = group_by_day(foods)
    training_days = set(training) | {row[0] for row in exercises} | {day for workout in workouts if (day := workout_day(workout))}
    current_month = datetime.now(TZ).date().replace(day=1)
    lines = [
        intro("生活与健康总览"),
        f'<div class="health-hero"><span>个人健康与生活记录</span><strong>把散落的数据放在一起看</strong><small>Apple 健康 · {html.escape(SCALE_SOURCE_LABEL)} · Scriptable · Excel 饮食表</small></div>',
        "数值均有原始来源；缺失值留空，可疑距离只在明细中标注。",
        stat_cards([
            ("体重（体脂秤）", f"{weight:.1f} kg" if weight is not None else "暂无", scale_day),
            ("体脂率（体脂秤）", f"{fat:.2f}%" if fat is not None else "暂无", scale_day),
            ("步数", f"约 {steps:,.0f}" if steps is not None else "暂无", step_day or "暂无日期"),
            ("静息心率", f"{resting:.0f} 次/分" if resting is not None else "暂无", heart_day or "暂无日期"),
            ("已标记睡着", f"{hours:.2f} 小时" if hours is not None else "暂无", sleep_day or "暂无日期"),
            ("饮食记录", f"{len(foods)} 条", f"最新 {food_day}" if food_day else "暂无日期"),
        ]),
        "## 快速进入",
        nav_callout(),
        f"体脂秤同步：{scale_status(scale_synced)}；最新测量日期：{scale_day}。",
        "## 最近 30 天健康趋势",
        "> [!info] 图表说明\n> 步数显示导出每日值的约数；心率使用静息心率；睡眠按醒来日汇总原始“睡着”片段并扣除重叠。距离数据待核对，未进入图表。",
        "### 步数",
        chart_block("bar", "步数", health_index.get("Steps", {}), "步"),
        "### 静息心率",
        chart_block("line", "静息心率", health_index.get("Resting Heart Rate", {}), "次/分"),
        "### 睡着片段合计",
        chart_block("bar", "睡着时长", sleep_hours, "小时"),
        "### 体重",
        chart_block("line", "体重", scale_series(scale_records, "weight"), "kg"),
        "## 饮食与训练日历",
        "日期格可进入当天专题记录；颜色代表饮食条目数。[[饮食训练日历|查看全部月份 →]]",
        calendar_table(current_month, food_days, training_days) if current_month else "暂无记录。",
        "## 数据范围",
        table(["来源", "已整理"], [
            ("Apple 健康", f"{metric_count} 条数值、{sleep_count} 条睡眠片段、{len(workouts)} 条体能训练"),
            ("饮食表", f"{len(foods)} 条饮食、{len(exercises)} 条运动"),
            ("Scriptable", f"{sum(len(x) for x in daily.values())} 条每日事项、{sum(len(x) for x in training.values())} 条健身记录"),
            (SCALE_SOURCE_LABEL, f"{len(scale_records)} 条测量记录"),
        ]),
        f"健康数据从 {min(days) if days else '暂无'} 起，最新日期 {max(days) if days else '暂无'}；已汇入 {len(health_files)} 份导出。历史范围取决于已导入文件。",
        f"饮食来源：{book.name}；最新记录 {food_day or '暂无'}。计划每天 {SCHEDULE_TIME:%H:%M} 更新。",
    ]
    return "\n\n".join(lines)


def write_page(name, body):
    VAULT.mkdir(parents=True, exist_ok=True)
    path = VAULT / f"{name}.md"
    tail = "\n\n## 个人笔记\n\n"
    if path.exists():
        existing = path.read_text(encoding="utf-8-sig")
        if END in existing:
            tail = existing.split(END, 1)[1]
        else:
            raise RuntimeError(f"目标页面已存在且不是本脚本生成，已停止以保护内容：{path}")
    frontmatter = "---\ncssclasses:\n  - health-space\n---\n\n"
    content = frontmatter + BEGIN + "\n\n" + body.strip() + "\n\n" + END + tail
    temp = path.with_suffix(".md.tmp")
    temp.write_text(content, encoding="utf-8")
    temp.replace(path)


def main():
    metrics, categories, workouts, health_files = health_data()
    daily, training = scriptable_data()
    foods, exercises, notes, profile, book = workbook_data()
    if not health_files:
        raise RuntimeError("尚未找到可读取的健康导出文件；未更新页面。")
    scale_synced = sync_scale()
    scale_records = scale_data()
    health_body, metric_count, days, health_index = metrics_page(metrics, categories, health_files)
    sleep_body, sleep_count, sleep_hours = sleep_page(categories)
    fitness_body, training_days = fitness_page(training, exercises, workouts)
    pages = {
        "生活与健康总览": overview_page(metric_count, days, health_index, sleep_hours, sleep_count, foods, exercises, daily, training, workouts, health_files, book, scale_records, scale_synced),
        "体脂秤": scale_page(scale_records, scale_synced),
        "健康指标": health_body,
        "睡眠记录": sleep_body,
        "健身与运动": fitness_body,
        "饮食记录": food_page(foods, notes, profile, book),
        "饮食训练日历": calendar_page(foods, training_days),
        "每日事项": tasks_page(daily),
    }
    for name, body in pages.items():
        write_page(name, body)
    print(f"完成：{len(pages)} 个页面，健康 {metric_count} 条，体脂秤 {len(scale_records)} 条，饮食 {len(foods)} 条，每日事项 {sum(map(len, daily.values()))} 条。")
    return scale_synced


def scheduled_run_needed(status_path, now=None):
    now = now or datetime.now(TZ)
    if now.time() < SCHEDULE_TIME:
        return False
    try:
        previous = read_json(status_path)
        previous_time = datetime.fromisoformat(previous["time"])
        if previous.get("status") == "ok" and previous_time.date() == now.date() and previous_time.time() >= SCHEDULE_TIME:
            return False
    except (OSError, ValueError, KeyError, TypeError):
        pass
    return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="将本地健康与生活数据整理为 Obsidian 页面")
    parser.add_argument("--config", default="config.json", help="本地配置文件路径")
    parser.add_argument("--scheduled", action="store_true", help="只在计划时间之后且今天未成功更新时运行")
    args = parser.parse_args()
    configure(args.config)
    status_path = ARCHIVE.parent / "last_run.json"
    status_path.parent.mkdir(parents=True, exist_ok=True)
    if args.scheduled and not scheduled_run_needed(status_path):
        sys.exit(0)
    try:
        scale_synced = main()
    except Exception as exc:
        status_path.write_text(json.dumps({"time": stamp(), "status": "error", "message": str(exc)}, ensure_ascii=False, indent=2), encoding="utf-8")
        raise
    else:
        status_path.write_text(json.dumps({"time": stamp(), "status": "partial" if scale_synced is False else "ok", "scale_sync": "fresh" if scale_synced else "cached" if scale_synced is False else "disabled"}, ensure_ascii=False, indent=2), encoding="utf-8")
